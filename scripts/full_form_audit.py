import sys
import os
import json
import urllib.request
import urllib.error
import time
import openpyxl
from concurrent.futures import ThreadPoolExecutor, as_completed

sys.stdout.reconfigure(encoding='utf-8')

BLOCKFROST_BASE = "https://cardano-preprod.blockfrost.io/api/v0"
project_id = "preprodQyce8v341ddratwaA7Vmq5yEVtDW5dQa"

excel_path = os.path.join(os.path.dirname(__file__), "..", "docs", "stellarvault-feedback.xlsx")
wb = openpyxl.load_workbook(excel_path)
sheet = wb.active

entries = []
for r in range(2, sheet.max_row + 1):
    wallet_val = sheet.cell(r, 5).value
    tx_val = sheet.cell(r, 6).value
    if wallet_val or tx_val or sheet.cell(r, 2).value:
        entries.append({
            "row": r,
            "time": str(sheet.cell(r, 1).value),
            "name": str(sheet.cell(r, 2).value or ""),
            "user_id": str(sheet.cell(r, 3).value or ""),
            "email": str(sheet.cell(r, 4).value or ""),
            "wallet": str(wallet_val).strip() if wallet_val else "",
            "tx_hash_submitted": str(tx_val).strip() if tx_val else ""
        })

print(f"Total rows extracted from Excel: {len(entries)}")

# Load synthetic-users.json
with open(os.path.join(os.path.dirname(__file__), "..", "docs", "synthetic-users.json"), "r", encoding="utf-8") as f:
    synth_data = json.load(f)

synth_map = {w["address"].strip(): w for w in synth_data["wallets"]}

def call_bf(endpoint):
    url = f"{BLOCKFROST_BASE}{endpoint}"
    req = urllib.request.Request(url, headers={"project_id": project_id})
    try:
        with urllib.request.urlopen(req, timeout=10) as resp:
            return resp.status, json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return e.code, None
    except Exception as e:
        return 0, str(e)

def verify_entry(entry):
    wallet = entry["wallet"]
    tx_sub = entry["tx_hash_submitted"]
    
    # 1. Check wallet onchain
    w_status, w_data = call_bf(f"/addresses/{wallet}")
    
    # 2. Check wallet real txs
    addr_txs = []
    if w_status == 200:
        txs_st, txs_dt = call_bf(f"/addresses/{wallet}/transactions?count=10")
        if txs_st == 200 and isinstance(txs_dt, list):
            addr_txs = [t["tx_hash"] for t in txs_dt]
            
    # 3. Check submitted tx hash
    tx_sub_status, tx_sub_data = (None, None)
    if tx_sub:
        tx_sub_status, tx_sub_data = call_bf(f"/txs/{tx_sub}")
        
    # Match with synthetic data
    synth_entry = synth_map.get(wallet)
    synth_lock_tx = synth_entry.get("lockTxHash") if synth_entry else None
    
    lovelace = 0
    if w_status == 200 and w_data and "amount" in w_data:
        for amt in w_data["amount"]:
            if amt.get("unit") == "lovelace":
                lovelace = int(amt.get("quantity", 0))
    
    return {
        "row": entry["row"],
        "user_id": entry["user_id"],
        "name": entry["name"],
        "email": entry["email"],
        "wallet": wallet,
        "wallet_is_valid_format": wallet.startswith("addr_test1"),
        "wallet_exists_onchain": (w_status == 200),
        "wallet_balance_ada": lovelace / 1_000_000,
        "wallet_onchain_tx_count": len(addr_txs),
        "wallet_actual_tx_hashes": addr_txs,
        "tx_hash_submitted": tx_sub,
        "tx_submitted_exists_onchain": (tx_sub_status == 200),
        "synth_dataset_match": bool(synth_entry),
        "synth_lock_tx": synth_lock_tx,
        "synth_lock_tx_in_actual_txs": (synth_lock_tx in addr_txs) if synth_lock_tx else False
    }

print("Starting parallel verification via Blockfrost...")
results = []
with ThreadPoolExecutor(max_workers=10) as executor:
    futures = {executor.submit(verify_entry, entry): entry for entry in entries}
    for f in as_completed(futures):
        res = f.result()
        results.append(res)
        print(f"Verified Row {res['row']:02d} ({res['user_id']}): Wallet onchain={res['wallet_exists_onchain']}, Actual TXs={res['wallet_onchain_tx_count']}, Submitted TX valid={res['tx_submitted_exists_onchain']}", flush=True)

# Sort results by row
results.sort(key=lambda x: x["row"])

out_path = os.path.join(os.path.dirname(__file__), "..", "docs", "form-verification-full.json")
with open(out_path, "w", encoding="utf-8") as f:
    json.dump(results, f, ensure_ascii=False, indent=2)

print("\n" + "="*60)
print("COMPREHENSIVE AUDIT SUMMARY:")
print("="*60)
total = len(results)
wallets_valid_fmt = sum(1 for r in results if r["wallet_is_valid_format"])
wallets_onchain = sum(1 for r in results if r["wallet_exists_onchain"])
wallets_with_txs = sum(1 for r in results if r["wallet_onchain_tx_count"] > 0)
submitted_txs_onchain = sum(1 for r in results if r["tx_submitted_exists_onchain"])
synth_match = sum(1 for r in results if r["synth_dataset_match"])
synth_lock_confirmed = sum(1 for r in results if r["synth_lock_tx_in_actual_txs"])

unique_wallets = set(r["wallet"] for r in results)

print(f"1. Total form responses (rows): {total}")
print(f"2. Unique wallets: {len(unique_wallets)}")
print(f"3. Valid Cardano address format: {wallets_valid_fmt}/{total} ({wallets_valid_fmt/total*100:.1f}%)")
print(f"4. Wallets REAL / active on Cardano Preprod: {wallets_onchain}/{total} ({wallets_onchain/total*100:.1f}%)")
print(f"5. Wallets with real onchain transactions: {wallets_with_txs}/{total} ({wallets_with_txs/total*100:.1f}%)")
print(f"6. Submitted TX hashes valid onchain: {submitted_txs_onchain}/{total} ({submitted_txs_onchain/total*100:.1f}%)")
print(f"7. Wallets matched with synthetic dataset: {synth_match}/{total}")
print(f"8. Synthetic lock tx confirmed onchain: {synth_lock_confirmed}/{total}")

# Check for duplicate wallet entries
if len(unique_wallets) < total:
    from collections import Counter
    counts = Counter(r["wallet"] for r in results)
    print("\nDuplicate wallet analysis:")
    for w, c in counts.items():
        if c > 1:
            dup_rows = [r["row"] for r in results if r["wallet"] == w]
            dup_uids = [r["user_id"] for r in results if r["wallet"] == w]
            print(f"  - Wallet {w[:25]}... appears {c} times in Rows {dup_rows} (User IDs: {dup_uids})")
