import sys
import os
import json
import urllib.request
import urllib.error
import time
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

BLOCKFROST_BASE = "https://cardano-preprod.blockfrost.io/api/v0"

# Read API key from scripts/.env
project_id = None
env_path = os.path.join(os.path.dirname(__file__), ".env")
if os.path.exists(env_path):
    with open(env_path, "r", encoding="utf-8") as f:
        for line in f:
            if line.startswith("BLOCKFROST_PROJECT_ID="):
                project_id = line.strip().split("=", 1)[1].strip('"\'')

if not project_id:
    print("ERROR: BLOCKFROST_PROJECT_ID not found in scripts/.env")
    sys.exit(1)

print(f"Using Blockfrost Project ID: {project_id[:10]}...")

# Load Excel
excel_path = os.path.join(os.path.dirname(__file__), "..", "docs", "stellarvault-feedback.xlsx")
wb = openpyxl.load_workbook(excel_path)
sheet = wb.active

print(f"Sheet title: {sheet.title}, Total rows: {sheet.max_row}, Total cols: {sheet.max_column}")

headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
print(f"Headers: {headers}")

entries = []
for r in range(2, sheet.max_row + 1):
    row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
    if any(row_vals):
        entries.append({
            "row": r,
            "time": sheet.cell(r, 1).value,
            "name": sheet.cell(r, 2).value,
            "user_id": sheet.cell(r, 3).value,
            "email": sheet.cell(r, 4).value,
            "wallet": str(sheet.cell(r, 5).value).strip() if sheet.cell(r, 5).value else "",
            "tx_hash": str(sheet.cell(r, 6).value).strip() if sheet.cell(r, 6).value else ""
        })

print(f"Total valid entries loaded: {len(entries)}")

def call_blockfrost(endpoint):
    url = f"{BLOCKFROST_BASE}{endpoint}"
    req = urllib.request.Request(url, headers={"project_id": project_id})
    try:
        with urllib.request.urlopen(req, timeout=10) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        return e.code, None
    except Exception as e:
        return 0, str(e)

results = []
for idx, entry in enumerate(entries, 1):
    wallet = entry["wallet"]
    tx_hash = entry["tx_hash"]
    
    wallet_status, wallet_data = call_blockfrost(f"/addresses/{wallet}")
    tx_status, tx_data = call_blockfrost(f"/txs/{tx_hash}") if tx_hash else (None, None)
    
    # Check if wallet has txs
    tx_count = 0
    if wallet_status == 200 and wallet_data:
        # Check tx count from address data
        # /addresses/{address}/transactions
        txs_status, txs_data = call_blockfrost(f"/addresses/{wallet}/transactions?count=10")
        if txs_status == 200 and isinstance(txs_data, list):
            tx_count = len(txs_data)

    res_item = {
        "index": idx,
        "row": entry["row"],
        "user_id": entry["user_id"],
        "name": entry["name"],
        "wallet": wallet,
        "wallet_valid_onchain": (wallet_status == 200),
        "wallet_status_code": wallet_status,
        "wallet_amount": wallet_data.get("amount") if (wallet_status == 200 and wallet_data) else None,
        "wallet_tx_count_sample": tx_count,
        "tx_hash": tx_hash,
        "tx_valid_onchain": (tx_status == 200),
        "tx_status_code": tx_status,
        "tx_block_height": tx_data.get("block_height") if (tx_status == 200 and tx_data) else None,
        "tx_block_time": tx_data.get("block_time") if (tx_status == 200 and tx_data) else None,
    }
    results.append(res_item)
    
    print(f"[{idx:02d}/{len(entries)}] {entry['user_id']} | Wallet onchain: {'YES' if wallet_status == 200 else 'NO (' + str(wallet_status) + ')'} (txs: {tx_count}) | TX onchain: {'YES' if tx_status == 200 else 'NO (' + str(tx_status) + ')'}")
    time.sleep(0.05) # avoid rate limits

# Save full results to json
out_path = os.path.join(os.path.dirname(__file__), "..", "docs", "form-verification-results.json")
with open(out_path, "w", encoding="utf-8") as f:
    json.dump(results, f, ensure_ascii=False, indent=2)

print(f"\nVerification finished! Output written to {out_path}")

# Summary statistics
wallets_valid = sum(1 for r in results if r["wallet_valid_onchain"])
txs_valid = sum(1 for r in results if r["tx_valid_onchain"])
print(f"SUMMARY:")
print(f"- Total entries checked: {len(results)}")
print(f"- Wallets found on Preprod onchain: {wallets_valid}/{len(results)}")
print(f"- Transaction hashes found on Preprod onchain: {txs_valid}/{len(results)}")

# Check duplicates or anomalies
unique_wallets = set(r["wallet"] for r in results)
unique_txs = set(r["tx_hash"] for r in results)
print(f"- Unique wallets: {len(unique_wallets)}")
print(f"- Unique tx hashes: {len(unique_txs)}")
